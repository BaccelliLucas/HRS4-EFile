import requests
from openpyxl import load_workbook
from time import sleep
import os
import shutil
from zipfile import BadZipFile, ZipFile
from datetime import date
from tkinter import filedialog
import time
from PIL import Image
from pypdf import PdfWriter

class Functions:

    def __init__(self, parent):
        self.parent = parent

        self.collaborators_file_path = None
        self.documents_file_path = None
        self.benefits_file_path = None

    def get_zip(self, url, headers, download_path, name:str):
        response = requests.get(url, headers=headers)
        if response.status_code == 200:
            os.makedirs(download_path, exist_ok=True)
            with open(os.path.join(download_path, f'{name}.zip'), "wb") as f:
                f.write(response.content)
            print("Zip salvo com sucesso!")
        else:
            return None
        
    def get_pdf(self, url, headers, download_path, name:str):
        response = requests.get(url, headers=headers)
        if response.status_code == 200:
            os.makedirs(download_path, exist_ok=True)
            with open(os.path.join(download_path, f'{name}.pdf'), "wb") as f:
                f.write(response.content)
            print("PDF salvo com sucesso!")
        else:
            return None

    def get_names(self, ws):
        list_names = []
        for index, col in enumerate(ws[1]):
            if col.value == "NOME":
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    list_names.append(row[index].value)
                return list_names

    def get_cpfs(self, ws):
        list_cpf = []
        for index, col in enumerate(ws[1]):
            if str(col.value).upper() == "CPF":
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    list_cpf.append(row[index].value)
                return list_cpf

    def get_collaborators_list(self, worksheet_name: object):
        list_sites = []
        list_edvs = []
        list_cpfs = []
        list_names = []
        collaborator_list = []
        for index, col in enumerate(worksheet_name[1]):
            if col.value == "PLANTA":
                for row in worksheet_name.iter_rows(min_row=2, max_row=worksheet_name.max_row):
                    list_sites.append(row[index].value)
            if col.value == "EDV":
                for row in worksheet_name.iter_rows(min_row=2, max_row=worksheet_name.max_row):
                    list_edvs.append(row[index].value)
            if col.value == "CPF":
                for row in worksheet_name.iter_rows(min_row=2, max_row=worksheet_name.max_row):
                    list_cpfs.append(row[index].value)
            if col.value == "NOME":
                for row in worksheet_name.iter_rows(min_row=2, max_row=worksheet_name.max_row):
                    list_names.append(row[index].value)

        for index, site in enumerate(list_sites):
            collaborator = {
                "site": site, 
                "edv": list_edvs[index],
                "cpf": list_cpfs[index],
                "name": list_names[index]
            }
            collaborator_list.append(collaborator)

        return collaborator_list 

    def get_edv_by_name(self, collaborator_list, collaborator_name: str):
        for _ in collaborator_list:
            col_name = _['name']
            col_edv = _['edv']
            if str(col_name).upper() == collaborator_name.upper():
                return col_edv
        return None

    def get_edv_by_cpf(self, collaborator_list, collaborator_cpf: str):
        for _ in collaborator_list:
            col_cpf = _['cpf']
            col_edv = _['edv']
            if str(col_cpf).upper() == collaborator_cpf.upper():
                return col_edv
        return None

    def rename_file_pdf(self, download_path, contains, collaborator_edv, code):
        while len(os.listdir(download_path)) == 0:
            sleep(1)
        file_folder_list = os.listdir(download_path)
        
        for file in file_folder_list:
            if contains in file:
                old_file = os.path.join(download_path, file)
                new_file = os.path.join(download_path, f"{collaborator_edv}_{code}.pdf")
                os.rename(old_file, new_file)

    def rename_file(self, download_path, contains, code):
        file_folder_list = os.listdir(download_path)
        
        for file in file_folder_list:
            if contains in file:
                old_file = os.path.join(download_path, file)
                new_file = os.path.join(download_path, f"{file}_{code}.pdf")
                os.rename(old_file, new_file)

    def rename_all(self, download_path, collaborator_edv):
        count_int = 0
        file_folder_list = os.listdir(download_path)
        valid_exts = ('.pdf', '.jpg', '.jpeg', '.png', '.heic', '.webp')
        
        for file in file_folder_list:
            if not file.lower().endswith(valid_exts):
                continue 

            count_int += 1
            formatted_number = f"{count_int:03d}"
            
            old_file = os.path.join(download_path, file)
            _, ext = os.path.splitext(file)
            new_file = os.path.join(download_path, f"{collaborator_edv}_{formatted_number}{ext.lower()}")
            
            os.rename(old_file, new_file)

    def delete_all_files(self, folder_path):
        if len(os.listdir(folder_path)) > 0:
            for folder in os.listdir(folder_path):
                all_folder_path = os.path.join(folder_path, folder)
                if os.path.isdir(all_folder_path):
                    for file in os.listdir(all_folder_path):
                        os.remove(os.path.join(all_folder_path, file))

    def move_files_to_BOT(self, current_folder, destination_path):
        if not os.path.isdir(current_folder):
            raise Exception(f'{current_folder} is not a folder!')
        
        valid_exts = ('.pdf', '.jpg', '.jpeg', '.png', '.heic', '.webp')
        
        for file in os.listdir(current_folder):
            if not file.lower().endswith(valid_exts): 
                continue
            
            file_path = os.path.join(current_folder, file)
            shutil.move(file_path, destination_path)

    def move_file_to_BOT(self, current_folder, file):
        if not os.path.isdir(current_folder):
            raise Exception(f'{current_folder} is not a folder!')
        
        file_path = os.path.join(current_folder, file)
        dst = r"C:\Users\pag4jvl\Documents\Automation\efile_selenium\pasta_bot_fake"
        shutil.move(file_path, dst)

    def safe_extract(self, zip_path, extract_to):
        try:
            with ZipFile(zip_path, 'r') as zip_ref:
                for member in zip_ref.infolist():
                    original_name = member.filename
                    fixed_name = os.path.basename(original_name.rstrip())
                    destination_path = os.path.join(extract_to, fixed_name)
                    abs_extract_path = os.path.abspath(extract_to)
                    abs_destination_path = os.path.abspath(destination_path)

                    if not abs_destination_path.startswith(abs_extract_path):
                        print(f"Aviso: Tentativa de extração para fora do diretório seguro: {original_name}")
                        continue

                    if member.is_dir():
                        os.makedirs(destination_path, exist_ok=True)
                        continue

                    os.makedirs(os.path.dirname(destination_path), exist_ok=True)

                    try:
                        with zip_ref.open(member) as source, open(destination_path, 'wb') as target:
                            target.write(source.read())
                    except Exception as e:
                        print(f"Erro ao extrair {original_name}: {e}")

        except FileNotFoundError:
            print(f"Erro: Arquivo ZIP não encontrado: {zip_path}")
        except BadZipFile:
            print(f"Erro: Arquivo ZIP corrompido ou inválido: {zip_path}")
        except Exception as e:
            print(f"Erro inesperado ao processar o arquivo ZIP: {e}")

    def correct_extension_by_signature(self, file_path):
        if not os.path.isfile(file_path):
            return file_path

        try:
            with open(file_path, 'rb') as f:
                header = f.read(12)
        except Exception:
            return file_path

        new_ext = None
        if header.startswith(b'%PDF'):
            new_ext = '.pdf'
        elif header.startswith(b'\x89PNG\r\n\x1a\n'):
            new_ext = '.png'
        elif header.startswith(b'\xff\xd8\xff'):
            new_ext = '.jpg'
        elif len(header) >= 12 and header[0:4] == b'RIFF' and header[8:12] == b'WEBP':
            new_ext = '.webp'
        elif len(header) >= 12 and header[4:8] == b'ftyp' and header[8:12] in (b'heic', b'mif1', b'heix', b'hevc'):
            new_ext = '.heic'

        if new_ext:
            base_name, current_ext = os.path.splitext(file_path)
            if current_ext.lower() != new_ext:
                new_file_path = base_name + new_ext
                if os.path.exists(new_file_path):
                    new_file_path = f"{base_name}_{int(time.time())}{new_ext}"
                os.rename(file_path, new_file_path)
                print(f"Extensão corrigida: {os.path.basename(file_path)} -> {new_ext}")
                return new_file_path
        return file_path

    def get_file_in_zip(self, download_folder):
        for file in os.listdir(download_folder):
            if file.endswith('.zip') or file.endswith(' .zip'):
                full_path = os.path.join(download_folder, file)
                self.safe_extract(full_path, download_folder)
                os.remove(full_path)

        for folder_zipped in os.listdir(download_folder):
            folder_path = os.path.join(download_folder, folder_zipped)
            if os.path.isdir(folder_path):
                for file in os.listdir(folder_path):
                    file_path = os.path.join(folder_path, file)
                    shutil.move(file_path, download_folder)
                try:
                    os.rmdir(folder_path)
                except OSError:
                    pass

        for file in os.listdir(download_folder):
            file_path = os.path.join(download_folder, file)
            if os.path.isfile(file_path):
                self.correct_extension_by_signature(file_path)

    def merge_files_to_pdf(self, folder_path, output_filename="contract.pdf"):
        files_to_merge = []
        valid_image_exts = ('.jpg', '.jpeg', '.png', '.webp')
        
        output_abs_path = os.path.join(folder_path, output_filename)

        for item in os.listdir(folder_path):
            file_path = os.path.join(folder_path, item)
            
            if not os.path.isfile(file_path) or file_path == output_abs_path:
                continue
            
            nome_base, extensao = os.path.splitext(item)
            ext_lower = extensao.lower()
            
            if ext_lower == '.pdf':
                files_to_merge.append(file_path)
                
            elif ext_lower in valid_image_exts:
                try:
                    image = Image.open(file_path).convert('RGB')
                    print(f"Convertendo imagem {item} para PDF temporário.")
                    
                    temp_pdf_name = f"{nome_base}.temp.pdf"
                    temp_pdf_path = os.path.join(folder_path, temp_pdf_name)
                    
                    image.save(temp_pdf_path)
                    files_to_merge.append(temp_pdf_path)
                    
                except Exception as e:
                    print(f"Erro ao converter imagem {item}: {e}")

        if not files_to_merge:
            print("Nenhum arquivo PDF ou imagem válido encontrado para mesclar.")
            return False

        print("\nIniciando mesclagem dos arquivos...")
        merger = PdfWriter()
        files_to_merge.sort() 
        
        for pdf_file in files_to_merge:
            try:
                merger.append(pdf_file)
            except Exception as e:
                print(f"Erro ao anexar o arquivo {os.path.basename(pdf_file)} à mesclagem: {e}")
                
        with open(output_abs_path, 'wb') as output_file:
            merger.write(output_file)
            
        print(f"✅ Concluído! Todos os arquivos foram mesclados em: {output_abs_path}")
        
        for item in os.listdir(folder_path):
            if item.endswith('.temp.pdf'):
                os.remove(os.path.join(folder_path, item))
        print("Arquivos temporários (.temp.pdf) limpos.\n")
        
        return True

    def fix_file_name(self, download_folder):
        for filePDF in os.listdir(download_folder):
            old_name = os.path.join(download_folder, filePDF)
            if not old_name.endswith('.pdf') and not old_name.endswith(' .pdf'):
                continue
            
            if filePDF.endswith(" .pdf"):
                nome_corrigido = filePDF.replace(" .pdf", ".pdf")
                caminho_novo = os.path.join(download_folder, nome_corrigido)
                os.rename(old_name, caminho_novo)

    def creating_txt_file(self, path_name, file_name, content):
        try:
            abs_path = os.path.join(path_name, file_name)
            with open(abs_path, 'w') as file_openned:
                file_openned.write(content)
            print(f'File {file_name} create successfully in {os.path.abspath(abs_path)}!')
        except IOError as e:
            print(f'Error creating file "{file_name}"')

    def appending_existing_file(self, path_name, file_name, content):
        try:
            abs_path = os.path.join(path_name, file_name)
            with open(abs_path, 'a') as file_oppened:
                file_oppened.write(content)
            print(f'File {file_name} append successfully!')
        except IOError as e:
            print(f'Error to append "{file_name}"!: {e}')

    def create_txt_in_file(self, collaborators_error_list, txt_folder):
        content = ''
        date_today = date.today().strftime("%d-%m-%Y")
        os.makedirs(txt_folder, exist_ok=True)
        file_name = "collaborators_error "+date_today+'.txt'
        path_abs = os.path.join(txt_folder, file_name)

        for collaborator in collaborators_error_list:
            content += f'{collaborator}\n'

        exit_file = False
        for file in os.listdir(txt_folder):
            if os.path.abspath(os.path.join(txt_folder, file)) == os.path.abspath(path_abs):
                self.appending_existing_file(txt_folder, file_name, content)
                exit_file = True

        if not exit_file:
            self.creating_txt_file(txt_folder, file_name, content)

    def verify_ok_spreadsheet(self, worksheet, row_index):
        col_index = -1
        filled_cell = False
        for index, col in enumerate(worksheet[1]):
            if col.value == "ENVIADO":
                col_index = col.column
                break

        if col_index != -1:
            row_to_update = worksheet.cell(row=row_index, column=col_index)
            if row_to_update.value == None:
                filled_cell = False
            else:
                filled_cell = True
        
        return filled_cell 

    def verify_ok_cell(self, worksheet, row_index):
        col_index = -1
        for index, col in enumerate(worksheet[1]):
            if col.value == "ENVIADO":
                col_index = col.column
                break

        if col_index != -1:
            row_to_update = worksheet.cell(row=row_index, column=col_index)
            if row_to_update.value == None:
                return False
            return True
        return False

    def insert_ok_spreadsheet(self, worksheet, workbook, worksheet_name, row_index):
        col_index = -1
        for index, col in enumerate(worksheet[1]):
            if col.value == "ENVIADO":
                col_index = col.column
                break

        if col_index != -1:
            row_to_update = worksheet.cell(row=row_index, column=col_index)
            if row_to_update.value == None:
                row_to_update.value = "OK"
                workbook.save(worksheet_name)
                print(f'A célula foi atualizada!')
            else:
                print('A célula já está preenchida!')

    def simulate_return_exception_return(self):
        try:
            raise requests.exceptions.ConnectionError("Simulando erro de conexão com a API")
        except requests.exceptions.ConnectionError as e:
            print("Erro de conexão com a API\n")
            return None
        except requests.exceptions.HTTPError as e:
            print('Erro de HTTP simulado\n')
            return None
        except Exception as e:
            print('Ocorreu um erro simulado inesparado!\n')

    def select_file(self, title):
        file_path_local = filedialog.askopenfilename(title=title, filetypes=[("Excel files", "*.xlsx *.xls")])
        if file_path_local:
            return file_path_local

    def select_folder(self, title):
        folder_path_local = filedialog.askdirectory(title=title)
        if folder_path_local:
            return folder_path_local